#!/usr/bin/env python3
"""
OpenPlag - open-source plagiarism + AI-content detector (zero per-query cost).

Pipeline
  ingest -> normalize -> segment -> plan queries -> discover sources (free
  backends) -> fetch/extract -> match (exact/near k-gram, paraphrase) ->
  aggregate -> AI-content signals -> report (with direct source links)

Required : Python 3.10+, `requests`
Optional : pymupdf | pypdf (PDF), python-docx (DOCX), antiword/libreoffice (DOC),
           trafilatura (better HTML extraction), pytesseract/ocrmypdf (OCR for scans),
           sentence-transformers (semantic paraphrase), torch + transformers (neural AI)
           -> see requirements.txt

Usage
  python openplag.py essay.pdf --json report.json
  python openplag.py --text "pasted text ..."
  python openplag.py essay.docx --html report.html --local-corpus ./refs
  python openplag.py --check-searx          # probe public SearXNG instances for JSON support
"""
from __future__ import annotations

import argparse
import hashlib
import html as htmllib
import json
import logging
import math
import os
import random
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import threading
import time
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from bisect import bisect_left
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field, asdict
from html.parser import HTMLParser
from pathlib import Path
from statistics import pstdev
from typing import Optional
from urllib import robotparser
from urllib.parse import parse_qs, quote, urlparse

try:
    import requests
except ImportError:  # offline/local-corpus mode still works
    requests = None

log = logging.getLogger("openplag")


# ============================================================================
# Configuration
# ============================================================================
@dataclass
class Config:
    # matching
    k: int = 5                      # shingle size (words)
    window: int = 4                 # winnowing window
    min_span_tokens: int = 8        # shortest reportable match
    max_gap: int = 4                # tolerated edit gap inside one match (tokens)
    min_density: float = 0.30       # share of k-grams that must survive in a span
    # paraphrase
    detect_paraphrase: bool = True
    lexical_threshold: float = 0.55
    use_semantic: bool = False      # needs sentence-transformers
    semantic_threshold: float = 0.82
    # discovery
    max_queries: int = 20
    results_per_query: int = 5
    max_sources: int = 30
    fetch_workers: int = 6
    exclude_domains: tuple = ()
    respect_robots: bool = True
    # exclusions
    exclude_quotes: bool = True
    exclude_references: bool = True
    exclude_citations: bool = True      # (Smith, 2020) and [1] markers don't count
    common_phrase_sources: int = 4      # short matches found in >= N sources = boilerplate
    common_phrase_max_tokens: int = 15
    store_max_df: int = 20              # ignore stored fingerprints seen in > N documents
    # ingestion
    ocr: bool = True                    # OCR scanned PDFs (pytesseract or ocrmypdf)
    # SearXNG failover pool. Best: your own instance(s) (see deploy/). Public instances are
    # opt-in (--public-searx) and best-effort: most disable JSON or rate-limit API-style clients.
    searx_instances: tuple = ()         # extra instances, tried in order after --searxng
    searx_discover: bool = True         # with --public-searx, fall back to searx.space listings
    # infra
    search_cache_hours: int = 24
    cache_max_mb: int = 500
    cache_dir: Path = field(default_factory=lambda: Path.home() / ".cache" / "openplag")
    cache_ttl_days: int = 7
    contact_email: str = ""         # polite-pool identification (OpenAlex etc.)
    user_agent: str = "OpenPlag/0.1 (open-source plagiarism checker)"


# ============================================================================
# Text utilities
# ============================================================================
STOP = frozenset("""a about above after again against all am an and any are as at be because been
before being below between both but by can could did do does doing down during each few for from
further had has have having he her here hers herself him himself his how i if in into is it its itself
just me more most my myself no nor not of off on once only or other our ours ourselves out over own
same she should so some such than that the their theirs them themselves then there these they this
those through to too under until up very was we were what when where which while who whom why will
with would you your yours yourself yourselves also may might shall upon within without however thus
""".split())

WORD_RE = re.compile(r"\w+(?:['\u2019]\w+)*", re.UNICODE)
_ZERO_WIDTH = dict.fromkeys(map(ord, "\u200b\u200c\u200d\u2060\ufeff\u00ad\u180e"), None)
# Cyrillic/Greek look-alikes used to evade checkers
_HOMOGLYPHS = {ord(a): b for a, b in zip(
    "\u0430\u0435\u043e\u0440\u0441\u0445\u0443\u0456\u0458\u0455\u04bb\u0410\u0412\u0415\u041a\u041c\u041d\u041e\u0420\u0421\u0422\u0425\u03bf\u0391\u0392\u0395\u0397\u0399\u039a\u039c\u039d\u039f\u03a1\u03a4\u03a7",
    "aeopcxyijshABEKMHOPCTXoABEHIKMNOPTX")}


def normalize_text(text: str) -> tuple[str, dict]:
    """Clean extraction artefacts and undo common evasion tricks."""
    stats = {"hidden_chars": sum(1 for ch in text if ord(ch) in _ZERO_WIDTH), "homoglyphs": 0}
    text = text.translate(_ZERO_WIDTH)
    text = unicodedata.normalize("NFKC", text)

    def fix(m: re.Match) -> str:
        w = m.group(0)
        if re.search(r"[A-Za-z]", w) and any(ord(c) in _HOMOGLYPHS for c in w):
            stats["homoglyphs"] += sum(ord(c) in _HOMOGLYPHS for c in w)
            return w.translate(_HOMOGLYPHS)
        return w

    text = WORD_RE.sub(fix, text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # PDF line-break hyphenation
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip(), stats


@dataclass
class Tokens:
    words: list
    starts: list
    ends: list


def tokenize(text: str) -> Tokens:
    w, s, e = [], [], []
    for m in WORD_RE.finditer(text):
        w.append(m.group(0).lower().replace("\u2019", "'"))
        s.append(m.start())
        e.append(m.end())
    return Tokens(w, s, e)


_ABBR = {"etc", "vs", "fig", "al", "dr", "mr", "mrs", "ms", "prof", "no", "eq", "cf", "inc", "ltd", "st", "jr", "sr"}
_SENT_END = re.compile(r'([.!?]+["\u201d\u2019)\]]*)\s+(?=[A-Z0-9"\u201c(\[])|\n\s*\n')


def split_sentences(text: str) -> list[tuple[int, int]]:
    spans, start = [], 0

    def add(a, b):
        seg = text[a:b]
        lead = len(seg) - len(seg.lstrip())
        trail = len(seg) - len(seg.rstrip())
        if seg.strip():
            spans.append((a + lead, b - trail))

    for m in _SENT_END.finditer(text):
        end = m.end(1) if m.group(1) else m.start()
        if m.group(1) and m.group(1).startswith("."):
            lw = re.search(r"([A-Za-z]+)$", text[start:m.start(1)])
            if lw and (lw.group(1).lower() in _ABBR or len(lw.group(1)) == 1):
                continue
        add(start, end)
        start = m.end()
    add(start, len(text))
    return spans


_QUOTE_RE = re.compile(r'\u201c([^\u201d]{25,}?)\u201d|"([^"\n]{25,}?)"')
_REF_RE = re.compile(r"^\s*(references|bibliography|works cited|reference list)\s*:?\s*$", re.I | re.M)
_CITE_RE = re.compile(
    r"\((?=[A-Z][^()]{2,119}\b(?:19|20)\d{2}[a-z]?\b)[^()]*\)"      # (Smith & Jones, 2020, p. 4)
    r"|\[\d+(?:\s*[,\u2013\-]\s*\d+)*\]")                           # [1]  [2, 5]  [3-7]


def _hash(gram) -> int:
    return int.from_bytes(hashlib.blake2b(" ".join(gram).encode(), digest_size=8).digest(), "big")


def kgram_hashes(words: list, k: int) -> list[int]:
    return [_hash(words[i:i + k]) for i in range(len(words) - k + 1)]


def winnow(hashes: list[int], w: int) -> dict[int, list[int]]:
    """Winnowing (Schleimer et al., MOSS). Guarantees detection of any shared run
    of >= w + k - 1 tokens while keeping ~2/(w+1) of the hashes. -> {hash: [positions]}"""
    fps: dict[int, list[int]] = defaultdict(list)
    if len(hashes) <= w:
        for i, h in enumerate(hashes):
            fps[h].append(i)
        return fps
    last = -1
    for i in range(len(hashes) - w + 1):
        win = hashes[i:i + w]
        m = min(win)
        j = i + (w - 1 - win[::-1].index(m))               # rightmost minimum
        if j != last:
            fps[m].append(j)
            last = j
    return fps


def content_count(words: list) -> int:
    return sum(1 for x in words if x not in STOP)


def _stem(w: str) -> str:
    for suf in ("ingly", "ing", "edly", "ed", "es", "ly", "s"):
        if len(w) > len(suf) + 3 and w.endswith(suf):
            return w[: -len(suf)]
    return w


def content_set(words: list) -> set:
    return {_stem(w) for w in words if w not in STOP and len(w) > 2}


# ============================================================================
# Document ingestion  (pdf / docx / doc / txt / html / raw text)
# ============================================================================
def _ocr_pdf(data: bytes, max_pages: int = 60) -> str:
    """OCR a scanned PDF: pytesseract (needs the tesseract binary) or the ocrmypdf CLI."""
    try:
        import io
        import fitz
        import pytesseract
        from PIL import Image
        out = []
        with fitz.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc):
                if i >= max_pages:
                    break
                img = Image.open(io.BytesIO(page.get_pixmap(dpi=200).tobytes("png")))
                out.append(pytesseract.image_to_string(img))
        return "\n".join(out)
    except ImportError:
        pass
    except Exception as e:                     # e.g. tesseract binary missing
        log.warning("pytesseract OCR failed: %s", e)
    exe = shutil.which("ocrmypdf")
    if exe:
        with tempfile.TemporaryDirectory() as td:
            src, side = Path(td) / "in.pdf", Path(td) / "out.txt"
            src.write_bytes(data)
            try:
                subprocess.run([exe, "--force-ocr", "-q", "--sidecar", str(side), str(src), str(Path(td) / "o.pdf")],
                               check=True, timeout=600, capture_output=True)
                return side.read_text(errors="replace")
            except Exception as e:
                log.warning("ocrmypdf failed: %s", e)
    log.warning("No OCR engine available: pip install pytesseract pillow (+ tesseract) or ocrmypdf")
    return ""


def pdf_to_text(data: bytes, max_pages: int = 300, ocr: bool = True) -> str:
    text = _pdf_native(data, max_pages)
    if ocr and len(text.strip()) < 100:
        log.info("PDF has no usable text layer - trying OCR")
        text = _ocr_pdf(data) or text
    return text


def _pdf_native(data: bytes, max_pages: int = 300) -> str:
    try:
        import fitz  # PyMuPDF
        with fitz.open(stream=data, filetype="pdf") as doc:
            return "\n".join(p.get_text() for i, p in enumerate(doc) if i < max_pages)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader
        import io
        rd = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in rd.pages[:max_pages])
    except ImportError:
        raise RuntimeError("Install pymupdf or pypdf to read PDFs")


def docx_to_text(path: str) -> str:
    try:
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception:  # python-docx missing or file odd: read the XML directly
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf8", "replace")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", "\t", xml)
        return htmllib.unescape(re.sub(r"<[^>]+>", "", xml))


def doc_to_text(path: str) -> str:
    """Legacy .doc via antiword or headless LibreOffice."""
    if shutil.which("antiword"):
        return subprocess.run(["antiword", path], capture_output=True, text=True, check=True).stdout
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        with tempfile.TemporaryDirectory() as td:
            subprocess.run([office, "--headless", "--convert-to", "txt:Text", "--outdir", td, path],
                           capture_output=True, check=True, timeout=120)
            return (Path(td) / (Path(path).stem + ".txt")).read_text(errors="replace")
    raise RuntimeError("Reading .doc needs antiword or LibreOffice installed")


class _TextExtractor(HTMLParser):
    SKIP = {"script", "style", "noscript", "nav", "footer", "aside", "form", "svg", "iframe", "head"}
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article", "blockquote", "h1", "h2", "h3",
             "h4", "h5", "h6", "title", "sec", "abstract"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.out, self.depth = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.depth += 1
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.depth = max(0, self.depth - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if self.depth == 0:
            self.out.append(data)


def html_to_text(markup: str) -> str:
    try:
        import trafilatura
        t = trafilatura.extract(markup, include_comments=False, include_tables=False)
        if t and len(t) > 200:
            return t
    except ImportError:
        pass
    p = _TextExtractor()
    try:
        p.feed(markup)
    except Exception:
        pass
    return "".join(p.out)


def extract_text(path: str, ocr: bool = True) -> str:
    """Extract plain text from a document on disk."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        text = pdf_to_text(p.read_bytes(), ocr=ocr)
        if len(text.strip()) < 50:
            log.warning("PDF has almost no text layer - it may be scanned; OCR (ocrmypdf) is needed")
        return text
    if ext == ".docx":
        return docx_to_text(str(p))
    if ext == ".doc":
        return doc_to_text(str(p))
    if ext in (".html", ".htm"):
        return html_to_text(p.read_text(errors="replace"))
    return p.read_text(errors="replace")            # .txt, .md, anything else textual


# ============================================================================
# HTTP (polite: per-host throttle, robots.txt, size cap) + disk cache
# ============================================================================
@dataclass
class Resp:
    status: int                      # 0 = network error, -1 = blocked by robots.txt
    data: bytes = b""
    ctype: str = ""
    enc: Optional[str] = None
    retry_after: float = 0.0


class Http:
    INTERVAL = {"html.duckduckgo.com": 2.0, "export.arxiv.org": 3.0, "api.openalex.org": 0.15,
                "www.ebi.ac.uk": 0.2, "en.wikipedia.org": 0.1, "api.crossref.org": 0.1,
                "api.semanticscholar.org": 1.5, "api.core.ac.uk": 1.0,
                "api2.marginalia-search.com": 2.0, "archive.org": 1.0}
    RETRY_STATUS = {429, 500, 502, 503, 504}

    def __init__(self, cfg: Config):
        self.cfg = cfg
        self.s = requests.Session() if requests else None
        if self.s:
            ua = cfg.user_agent + (f" mailto:{cfg.contact_email}" if cfg.contact_email else "")
            self.s.headers["User-Agent"] = ua
        self._last: dict = {}
        self._lock = threading.Lock()
        self._robots: dict = {}

    def _wait(self, host: str, gap: Optional[float] = None):
        gap = self.INTERVAL.get(host, 0.5) if gap is None else gap
        with self._lock:
            now = time.monotonic()
            slot = max(now, self._last.get(host, 0) + gap)
            self._last[host] = slot
        if slot > now:
            time.sleep(slot - now)

    def allowed(self, url: str) -> bool:
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        rp = self._robots.get(base)
        if rp is None:
            rp = robotparser.RobotFileParser()
            try:
                r = self.s.get(base + "/robots.txt", timeout=8)
                rp.parse(r.text.splitlines() if r.status_code == 200 else [])
            except Exception:
                rp.parse([])
            self._robots[base] = rp
        return rp.can_fetch(self.cfg.user_agent, url)

    def request(self, url, params=None, headers=None, max_bytes=8_000_000, timeout=15,
                robots=False, retries=2, gap=None) -> Resp:
        """GET with per-host throttling and exponential backoff (honours Retry-After)
        on 429/5xx and network errors. Never raises."""
        if not self.s:
            return Resp(0)
        if robots and self.cfg.respect_robots and not self.allowed(url):
            log.debug("robots.txt disallows %s", url)
            return Resp(-1)
        host = urlparse(url).hostname or ""
        resp = Resp(0)
        for attempt in range(retries + 1):
            self._wait(host, gap)
            try:
                r = self.s.get(url, params=params, headers=headers, timeout=timeout, stream=True)
                ra = str(r.headers.get("Retry-After", ""))
                buf = bytearray()
                if r.status_code == 200:
                    for chunk in r.iter_content(65536):
                        buf.extend(chunk)
                        if len(buf) > max_bytes:
                            break
                r.close()
                resp = Resp(r.status_code, bytes(buf), r.headers.get("Content-Type", ""), r.encoding,
                            float(ra) if ra.isdigit() else 0.0)
            except Exception as e:
                log.debug("GET %s failed: %s", url, e)
                resp = Resp(0)
            if resp.status == 200 or (resp.status and resp.status not in self.RETRY_STATUS):
                return resp
            if attempt < retries:
                time.sleep(min(20.0, resp.retry_after or 2 ** attempt) + random.uniform(0, 0.4))
        return resp

    def get(self, url, params=None, max_bytes=8_000_000, timeout=15, robots=False,
            headers=None, retries=2):
        """-> (bytes, content_type, encoding) on HTTP 200, else None"""
        r = self.request(url, params=params, headers=headers, max_bytes=max_bytes, timeout=timeout,
                         robots=robots, retries=retries)
        return (r.data, r.ctype, r.enc) if r.status == 200 else None


class DiskCache:
    """Tiny JSON file cache. Empty-string values are *negative* entries (1h TTL).
    Oldest files are pruned when the directory exceeds max_mb."""
    NEG_TTL = 3600

    def __init__(self, d: Path, ttl_days: int, max_mb: int = 500):
        self.d, self.ttl, self.max_bytes = Path(d), ttl_days * 86400, max_mb * 1_000_000
        try:
            self.d.mkdir(parents=True, exist_ok=True)
            self._prune()
        except OSError:
            pass

    def _prune(self):
        files = [(f.stat().st_mtime, f.stat().st_size, f) for f in self.d.glob("*.json")]
        total = sum(sz for _, sz, _ in files)
        if total <= self.max_bytes:
            return
        for _, sz, f in sorted(files):
            f.unlink(missing_ok=True)
            total -= sz
            if total <= self.max_bytes * 0.8:
                break

    def _p(self, key):
        return self.d / (hashlib.sha1(key.encode()).hexdigest() + ".json")

    def get(self, key, ttl: Optional[float] = None) -> Optional[str]:
        try:
            o = json.loads(self._p(key).read_text())
            limit = self.NEG_TTL if o["text"] == "" else (ttl or self.ttl)
            return o["text"] if time.time() - o["ts"] < limit else None
        except Exception:
            return None

    def put(self, key, text: str):
        try:
            self._p(key).write_text(json.dumps({"ts": time.time(), "text": text}))
        except OSError:
            pass

    def get_json(self, key, ttl: Optional[float] = None):
        t = self.get(key, ttl)
        try:
            return json.loads(t) if t else None
        except ValueError:
            return None

    def put_json(self, key, obj):
        self.put(key, json.dumps(obj))


# ============================================================================
# Source discovery - all free
# ============================================================================
@dataclass
class Source:
    url: str
    title: str = ""
    kind: str = "web"          # web | academic | preprint | encyclopedia | local | corpus
    text: str = ""
    fetch_url: str = ""        # where to download full text if different from url
    votes: int = 0


def _canon(url: str) -> str:
    p = urlparse(url)
    q = "&".join(x for x in p.query.split("&") if x and not x.startswith(("utm_", "fbclid", "gclid")))
    return f"{p.scheme}://{(p.hostname or '').lower()}{p.path.rstrip('/')}" + (f"?{q}" if q else "")


def _json(got):
    try:
        return json.loads(got[0].decode("utf-8", "replace")) if got else {}
    except Exception:
        return {}


class Backend:
    name = "base"

    def __init__(self, http: Http, cfg: Config):
        self.http, self.cfg = http, cfg

    def search(self, phrase: str, n: int) -> list[Source]:
        raise NotImplementedError


class DuckDuckGo(Backend):
    """Free HTML endpoint (scraped; may rate-limit). Prefer SearXNG for production."""
    name = "ddg"

    def search(self, phrase, n):
        got = self.http.get("https://html.duckduckgo.com/html/", params={"q": f'"{phrase}"'})
        if not got:
            return []
        page = got[0].decode(got[2] or "utf-8", "replace")
        out = []
        for m in re.finditer(r'<a\b([^>]*class="[^"]*result__a[^"]*"[^>]*)>(.*?)</a>', page, re.S):
            href = re.search(r'href="([^"]+)"', m.group(1))
            if not href:
                continue
            url = htmllib.unescape(href.group(1))
            if url.startswith("//"):
                url = "https:" + url
            qs = parse_qs(urlparse(url).query)
            if "uddg" in qs:
                url = qs["uddg"][0]
            if not url.startswith("http") or "duckduckgo.com" in url:
                continue
            title = htmllib.unescape(re.sub(r"<[^>]+>", "", m.group(2))).strip()
            out.append(Source(url=url, title=title, kind="web"))
            if len(out) >= n:
                break
        return out


class SearXNGPool(Backend):
    """
    Failover pool of SearXNG instances (your own first, then public ones).
    Instances are tried in order; one that returns 403 (JSON disabled), non-JSON,
    429/503, errors, or only blocked engines is put on cooldown and the next is used.
    If every listed instance is down, instances listed on searx.space are tried.
    Public instances change often - run `python openplag.py --check-searx` to see
    which currently serve JSON.
    """
    name = "searxng"
    DISCOVERY_URL = "https://searx.space/data/instances.json"
    LONG, SHORT = 86400, 3600          # cooldown seconds: JSON unsupported / rate-limited or down

    def __init__(self, http, cfg, instances=(), discover=False, cooldown: int = SHORT, cache=None):
        super().__init__(http, cfg)
        self.instances = list(dict.fromkeys(u.rstrip("/") for u in instances if u))
        self.discover, self.cooldown, self.cache = discover, cooldown, cache
        self._down: dict[str, float] = {}
        self._discovered = False
        self._lock = threading.Lock()

    # -- health ---------------------------------------------------------------
    def _usable(self, base: str) -> bool:
        if self._down.get(base, 0) > time.time():
            return False
        if self.cache:                               # health persists across runs
            d = self.cache.get_json(f"searxdown:{base}", ttl=self.LONG + 60)
            if d and d.get("until", 0) > time.time():
                self._down[base] = d["until"]
                return False
        return True

    def _mark(self, base: str, seconds: float, why: str):
        until = time.time() + seconds
        with self._lock:
            self._down[base] = until
        if self.cache:
            self.cache.put_json(f"searxdown:{base}", {"until": until, "why": why})
        log.info("SearXNG %s unavailable (%s) for %ds", base, why, seconds)

    # -- one request ------------------------------------------------------------
    def _query(self, base: str, phrase: str, n: int):
        r = self.http.request(base + "/search", params={"q": f'"{phrase}"', "format": "json"},
                              retries=0, timeout=12, gap=2.5)
        if r.status == 200:
            try:
                d = json.loads(r.data.decode("utf-8", "replace"))
            except ValueError:
                return None, "not_json"
            if not isinstance(d, dict) or "results" not in d:
                return None, "not_json"
            res = [Source(url=x["url"], title=x.get("title", ""), kind="web")
                   for x in d["results"] if x.get("url")]
            if not res and len(d.get("unresponsive_engines") or []) >= 3:
                return None, "engines_blocked"
            return res, "ok"
        if r.status == 403:
            return None, "json_disabled"
        if r.status in (429, 503):
            return None, "rate_limited"
        return None, "unreachable" if r.status == 0 else f"http_{r.status}"

    # -- discovery ----------------------------------------------------------------
    def discover_instances(self, limit: int = 8) -> list[str]:
        key = "searx:discovered"
        cached = self.cache.get_json(key, ttl=6 * 3600) if self.cache else None
        if cached:
            return cached
        r = self.http.request(self.DISCOVERY_URL, retries=1, timeout=30, max_bytes=30_000_000)
        if r.status != 200:
            return []
        try:
            data = json.loads(r.data.decode("utf-8", "replace")).get("instances", {})
        except ValueError:
            return []
        scored = []
        for url, info in data.items():
            if ".onion" in url or ".i2p" in url:
                continue
            if ((info or {}).get("http") or {}).get("status_code") != 200:
                continue
            t = ((info or {}).get("timing") or {}).get("search") or {}
            ok = t.get("success_percentage") or 0
            med = (t.get("all") or {}).get("median") or 9
            scored.append((-ok, med, url.rstrip("/")))
        found = [u for _, _, u in sorted(scored)[:limit]]
        if found and self.cache:
            self.cache.put_json(key, found)
        return found

    def _try_discover(self) -> bool:
        with self._lock:
            if not self.discover or self._discovered:
                return False
            self._discovered = True
        new = [u for u in self.discover_instances() if u not in self.instances]
        with self._lock:
            self.instances += new
        return bool(new)

    # -- API ----------------------------------------------------------------------
    def search(self, phrase, n):
        for attempt in range(2):
            for base in list(self.instances):
                if not self._usable(base):
                    continue
                res, why = self._query(base, phrase, n)
                if res is not None:
                    return res[:n]
                self._mark(base, self.LONG if why in ("json_disabled", "not_json") else self.cooldown, why)
            if attempt == 0 and not self._try_discover():
                break
        return []

    def probe(self, urls=None) -> list[dict]:
        """Check which instances currently serve JSON search results."""
        out = []
        for base in (urls or list(self.instances)):
            t = time.time()
            res, why = self._query(base, "photosynthesis chlorophyll", 3)
            out.append({"instance": base, "ok": bool(res), "reason": why if not res else "ok",
                        "results": len(res or []), "ms": int((time.time() - t) * 1000)})
        return out


class Crossref(Backend):
    """Crossref metadata (DOIs, abstracts when deposited). Low yield for body text."""
    name = "crossref"

    def search(self, phrase, n):
        params = {"query": phrase, "rows": n, "select": "DOI,title,URL,abstract,link"}
        if self.cfg.contact_email:
            params["mailto"] = self.cfg.contact_email
        d = _json(self.http.get("https://api.crossref.org/works", params=params))
        out = []
        for it in d.get("message", {}).get("items", []):
            abstract = htmllib.unescape(re.sub(r"<[^>]+>", " ", it.get("abstract") or "")).strip()
            pdf = next((l.get("URL", "") for l in it.get("link", []) if "pdf" in (l.get("content-type") or "")), "")
            url = it.get("URL") or (f"https://doi.org/{it['DOI']}" if it.get("DOI") else "")
            if url and (abstract or pdf):
                out.append(Source(url=url, title=(it.get("title") or [""])[0], kind="academic",
                                  text=abstract, fetch_url=pdf))
        return out


class SemanticScholar(Backend):
    """Free tier (shared, easily rate-limited). Set S2_API_KEY for a dedicated limit."""
    name = "semanticscholar"

    def search(self, phrase, n):
        key = os.environ.get("S2_API_KEY")
        d = _json(self.http.get(
            "https://api.semanticscholar.org/graph/v1/paper/search",
            params={"query": phrase, "limit": n, "fields": "title,abstract,url,openAccessPdf,externalIds"},
            headers={"x-api-key": key} if key else None))
        out = []
        for it in d.get("data", []):
            doi = (it.get("externalIds") or {}).get("DOI")
            url = f"https://doi.org/{doi}" if doi else it.get("url")
            pdf = (it.get("openAccessPdf") or {}).get("url") or ""
            if url and (it.get("abstract") or pdf):
                out.append(Source(url=url, title=it.get("title") or "", kind="academic",
                                  text=it.get("abstract") or "", fetch_url=pdf))
        return out


class CORE(Backend):
    """CORE open-access aggregator. Needs a free key: https://core.ac.uk/services/api -> CORE_API_KEY."""
    name = "core"

    def search(self, phrase, n):
        key = os.environ.get("CORE_API_KEY")
        if not key:
            return []
        d = _json(self.http.get("https://api.core.ac.uk/v3/search/works",
                                params={"q": f'"{phrase}"', "limit": n},
                                headers={"Authorization": f"Bearer {key}"}))
        out = []
        for it in d.get("results", []):
            doi = it.get("doi")
            url = f"https://doi.org/{doi}" if doi else (it.get("downloadUrl") or "")
            text = it.get("fullText") or it.get("abstract") or ""
            if url and (text or it.get("downloadUrl")):
                out.append(Source(url=url, title=it.get("title") or "", kind="academic", text=text,
                                  fetch_url="" if it.get("fullText") else (it.get("downloadUrl") or "")))
        return out


class Marginalia(Backend):
    """Independent/small-web index. 'public' key is shared and often overloaded (503)."""
    name = "marginalia"

    def search(self, phrase, n):
        key = os.environ.get("MARGINALIA_API_KEY", "public")
        d = _json(self.http.get("https://api2.marginalia-search.com/search",
                                params={"query": phrase, "count": n}, headers={"API-Key": key}, retries=1))
        return [Source(url=r["url"], title=r.get("title", ""), kind="web")
                for r in d.get("results", []) if r.get("url")]


class InternetArchive(Backend):
    """EXPERIMENTAL: advancedsearch matches item metadata, not page text, so yield is low.
    (archive.org's full-text-search API is still beta and changing.)"""
    name = "archive"

    def search(self, phrase, n):
        terms = [w for w in dict.fromkeys(phrase.split()) if w not in STOP and len(w) > 3][:5]
        if not terms:
            return []
        q = "mediatype:texts AND (" + " AND ".join(terms) + ")"
        d = _json(self.http.get("https://archive.org/advancedsearch.php", params=[
            ("q", q), ("fl[]", "identifier"), ("fl[]", "title"), ("rows", n), ("output", "json")]))
        out = []
        for it in d.get("response", {}).get("docs", []):
            i = it.get("identifier")
            if i:
                out.append(Source(url=f"https://archive.org/details/{i}", title=str(it.get("title", "")),
                                  kind="library", fetch_url=f"https://archive.org/download/{i}/{i}_djvu.txt"))
        return out


class Wikipedia(Backend):
    name = "wikipedia"

    def search(self, phrase, n):
        d = _json(self.http.get("https://en.wikipedia.org/w/api.php", params={
            "action": "query", "list": "search", "srsearch": f'"{phrase}"',
            "srlimit": n, "format": "json", "srprop": ""}))
        out = []
        for r in d.get("query", {}).get("search", []):
            t = r["title"]
            out.append(Source(url="https://en.wikipedia.org/wiki/" + quote(t.replace(" ", "_")),
                              title=t, kind="encyclopedia"))
        return out


class OpenAlex(Backend):
    """Free scholarly index. Gives abstracts and open-access PDF links."""
    name = "openalex"

    @staticmethod
    def _abstract(inv: Optional[dict]) -> str:
        if not inv:
            return ""
        pos = [(i, w) for w, idxs in inv.items() for i in idxs]
        return " ".join(w for _, w in sorted(pos))

    def search(self, phrase, n):
        params = {"search": phrase, "per-page": n,
                  "select": "id,doi,title,abstract_inverted_index,primary_location,best_oa_location"}
        if self.cfg.contact_email:
            params["mailto"] = self.cfg.contact_email
        out = []
        for w in _json(self.http.get("https://api.openalex.org/works", params=params)).get("results", []):
            loc = w.get("primary_location") or {}
            url = w.get("doi") or loc.get("landing_page_url") or w.get("id")
            pdf = (w.get("best_oa_location") or {}).get("pdf_url") or ""
            abstract = self._abstract(w.get("abstract_inverted_index"))
            if url and (abstract or pdf):
                out.append(Source(url=url, title=w.get("title") or "", kind="academic",
                                  text=abstract, fetch_url=pdf))
        return out


class EuropePMC(Backend):
    """Free full-text phrase search over open-access biomedical literature."""
    name = "europepmc"

    def search(self, phrase, n):
        d = _json(self.http.get("https://www.ebi.ac.uk/europepmc/webservices/rest/search", params={
            "query": f'"{phrase}"', "format": "json", "pageSize": n, "resultType": "lite"}))
        out = []
        for it in d.get("resultList", {}).get("result", []):
            pmcid = it.get("pmcid")
            if not pmcid:
                continue
            url = ("https://doi.org/" + it["doi"]) if it.get("doi") else \
                  f"https://europepmc.org/article/{it.get('source')}/{it.get('id')}"
            out.append(Source(url=url, title=it.get("title", ""), kind="academic",
                              fetch_url=f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/fullTextXML"))
        return out


class ArXiv(Backend):
    """arXiv indexes title/abstract metadata only - low yield for body text."""
    name = "arxiv"

    def search(self, phrase, n):
        got = self.http.get("http://export.arxiv.org/api/query",
                            params={"search_query": f'all:"{phrase}"', "max_results": n})
        if not got:
            return []
        ns = {"a": "http://www.w3.org/2005/Atom"}
        out = []
        try:
            for e in ET.fromstring(got[0]).findall("a:entry", ns):
                pdf = next((l.get("href") for l in e.findall("a:link", ns) if l.get("title") == "pdf"), "")
                out.append(Source(url=e.findtext("a:id", "", ns), title=(e.findtext("a:title", "", ns) or "").strip(),
                                  kind="preprint", text=e.findtext("a:summary", "", ns) or "", fetch_url=pdf))
        except ET.ParseError:
            pass
        return out


BACKENDS = {"ddg": DuckDuckGo, "wikipedia": Wikipedia, "openalex": OpenAlex, "europepmc": EuropePMC,
            "arxiv": ArXiv, "crossref": Crossref, "semanticscholar": SemanticScholar, "core": CORE,
            "marginalia": Marginalia, "archive": InternetArchive}
DEFAULT_BACKENDS = ("wikipedia", "openalex", "europepmc", "crossref")   # + SearXNG pool; core/S2 auto-on with keys


def build_backends(http, cfg, names=DEFAULT_BACKENDS, searxng_url: str = "",
                   public_searx: bool = False, cache=None) -> list[Backend]:
    names = list(names)
    for env, n in (("CORE_API_KEY", "core"), ("S2_API_KEY", "semanticscholar")):
        if os.environ.get(env) and n not in names:
            names.append(n)
    out = [BACKENDS[n](http, cfg) for n in names if n in BACKENDS]
    inst = ([searxng_url] if searxng_url else []) + (list(cfg.searx_instances) if public_searx else [])
    if inst or public_searx:
        out.append(SearXNGPool(http, cfg, inst, discover=public_searx and cfg.searx_discover, cache=cache))
    return out


# ============================================================================
# Persistent winnowed-fingerprint store (previous submissions / offline corpora)
# Stores ONLY hashes, never text -> privacy friendly, tiny, fast.
# ============================================================================
class FingerprintStore:
    def __init__(self, path: str = "openplag_store.db", cfg: Optional[Config] = None):
        self.cfg = cfg or Config()
        self.db = sqlite3.connect(path, check_same_thread=False)
        self.db.executescript("""
            CREATE TABLE IF NOT EXISTS docs(id TEXT PRIMARY KEY, title TEXT, url TEXT, kind TEXT);
            CREATE TABLE IF NOT EXISTS fps(h INTEGER, doc TEXT);
            CREATE INDEX IF NOT EXISTS fps_h ON fps(h);""")

    @staticmethod
    def _signed(h: int) -> int:                       # SQLite ints are signed 64-bit
        return h - (1 << 64) if h >= (1 << 63) else h

    def add(self, text: str, title: str = "", url: str = "", kind: str = "corpus") -> str:
        clean, _ = normalize_text(text)
        words = tokenize(clean).words
        doc_id = hashlib.sha1(" ".join(words).encode()).hexdigest()
        if self.db.execute("SELECT 1 FROM docs WHERE id=?", (doc_id,)).fetchone():
            return doc_id
        fps = winnow(kgram_hashes(words, self.cfg.k), self.cfg.window)
        with self.db:
            self.db.execute("INSERT INTO docs VALUES(?,?,?,?)", (doc_id, title, url, kind))
            self.db.executemany("INSERT INTO fps VALUES(?,?)", [(self._signed(h), doc_id) for h in fps])
        return doc_id

    def query(self, doc_fps: dict[int, list[int]]) -> dict[str, list[int]]:
        """-> {stored_doc_id: sorted submission token positions of shared fingerprints}.
        Fingerprints present in more than cfg.store_max_df stored documents are boilerplate and skipped."""
        by_hash: dict[int, set] = defaultdict(set)
        keys = list(doc_fps)
        for i in range(0, len(keys), 500):
            chunk = keys[i:i + 500]
            q = ",".join("?" * len(chunk))
            for h, d in self.db.execute(f"SELECT h, doc FROM fps WHERE h IN ({q})",
                                        [self._signed(x) for x in chunk]):
                by_hash[h if h >= 0 else h + (1 << 64)].add(d)
        hits: dict[str, set] = defaultdict(set)
        for h, docs in by_hash.items():
            if len(docs) > self.cfg.store_max_df:
                continue
            for d in docs:
                hits[d].update(doc_fps[h])
        return {d: sorted(p) for d, p in hits.items()}

    def meta(self, doc_id: str) -> tuple:
        return self.db.execute("SELECT title, url, kind FROM docs WHERE id=?", (doc_id,)).fetchone() or ("", "", "corpus")


# ============================================================================
# Prepared documents and matching engine
# ============================================================================
class Prepared:
    def __init__(self, text: str, k: int):
        self.text, self.k = text, k
        self.tok = tokenize(text)
        self.hashes = kgram_hashes(self.tok.words, k)
        self._index: Optional[dict] = None

    @property
    def index(self) -> dict:
        if self._index is None:
            idx: dict[int, list[int]] = {}
            for i, h in enumerate(self.hashes):
                lst = idx.setdefault(h, [])
                if len(lst) < 3:
                    lst.append(i)
            self._index = idx
        return self._index


@dataclass
class Raw:
    a: int                  # submission token range [a, b)
    b: int
    src: int                # index into sources list
    kind: str               # exact | near | paraphrase | corpus
    sim: float
    src_a: Optional[int] = None


def find_spans(doc: Prepared, src: Prepared, src_idx: int, cfg: Config) -> list[Raw]:
    """Exact / near-copy detection: shared k-grams merged with edit-gap tolerance."""
    k, words = cfg.k, doc.tok.words
    sidx = src.index
    hits = [(i, sidx[h][0]) for i, h in enumerate(doc.hashes) if h in sidx]
    if not hits:
        return []
    runs, cur = [], [hits[0]]
    for h in hits[1:]:
        if h[0] <= cur[-1][0] + k + cfg.max_gap:
            cur.append(h)
        else:
            runs.append(cur)
            cur = [h]
    runs.append(cur)
    out = []
    for r in runs:
        a, b = r[0][0], r[-1][0] + k
        possible = (b - a) - k + 1
        density = len(r) / possible
        if (b - a) < cfg.min_span_tokens or len(r) < 2 or density < cfg.min_density:
            continue
        if content_count(words[a:b]) < 3:
            continue
        out.append(Raw(a, b, src_idx, "exact" if density >= 0.97 else "near", round(min(density, 1.0), 3), r[0][1]))
    return out


def find_corpus_spans(doc: Prepared, positions: list[int], src_idx: int, cfg: Config) -> list[Raw]:
    """Spans from sparse winnowed hits (stored fingerprints, no source text)."""
    k, w, words = cfg.k, cfg.window, doc.tok.words
    out, cur = [], [positions[0]]
    for p in positions[1:] + [None]:
        if p is not None and p <= cur[-1] + w + cfg.max_gap:
            cur.append(p)
            continue
        a, b = cur[0], cur[-1] + k
        if len(cur) >= 2 and (b - a) >= cfg.min_span_tokens and content_count(words[a:b]) >= 3:
            expected = max(1.0, (b - a - k + 1) * 2 / (w + 1))
            out.append(Raw(a, b, src_idx, "corpus", round(min(1.0, len(cur) / expected), 3)))
        if p is not None:
            cur = [p]
    return out


def _sentence_token_ranges(doc: Prepared, sent_spans) -> list[tuple[int, int]]:
    st = doc.tok.starts
    return [(bisect_left(st, a), bisect_left(st, b)) for a, b in sent_spans]


def paraphrase_pass(doc: Prepared, sent_ranges, covered: list[bool], excluded: list[bool],
                    sources: list[Source], src_preps: dict, cfg: Config, embedder=None) -> list[Raw]:
    """Catch reworded copying on sentences not already matched verbatim."""
    todo = []
    for (a, b) in sent_ranges:
        if b - a < 7 or sum(excluded[a:b]) > (b - a) / 2:
            continue
        if sum(covered[a:b]) / (b - a) >= 0.4:
            continue
        cs = content_set(doc.tok.words[a:b])
        if len(cs) >= 5:
            todo.append((a, b, cs))
    if not todo:
        return []
    best: dict[int, Raw] = {}
    for si, sp in src_preps.items():
        text = sp.text
        s_spans = split_sentences(text)[:4000]
        s_sets, inv = [], defaultdict(list)
        for j, (x, y) in enumerate(s_spans):
            cs = content_set([m.group(0).lower() for m in WORD_RE.finditer(text[x:y])])
            s_sets.append(cs)
            for t in cs:
                inv[t].append(j)
        sem = None
        if embedder is not None and s_spans:
            import numpy as np
            s_emb = embedder.encode([text[x:y] for x, y in s_spans], normalize_embeddings=True, batch_size=64)
            d_emb = embedder.encode([" ".join(doc.tok.words[a:b]) for a, b, _ in todo],
                                    normalize_embeddings=True, batch_size=64)
            sem = d_emb @ s_emb.T
        for ti, (a, b, cs) in enumerate(todo):
            cnt = Counter(j for t in cs for j in inv.get(t, ()))
            cand = [(j, c) for j, c in cnt.most_common(5) if c >= 3]
            lex_best, lex_j = 0.0, -1
            for j, _ in cand:
                jac = len(cs & s_sets[j]) / len(cs | s_sets[j])
                if jac > lex_best:
                    lex_best, lex_j = jac, j
            score = 0.0
            if lex_best >= cfg.lexical_threshold:
                score = lex_best
            if sem is not None:
                j = int(sem[ti].argmax())
                c = float(sem[ti][j])
                lex_here = len(cs & s_sets[j]) / max(1, len(cs | s_sets[j]))
                if c >= cfg.semantic_threshold and lex_here >= 0.2:
                    score = max(score, c)
            if score and (a not in best or score > best[a].sim):
                best[a] = Raw(a, b, si, "paraphrase", round(score, 3))
    return list(best.values())


def plan_queries(doc: Prepared, sent_ranges, excluded: list[bool], n: int, win: int = 10) -> list[str]:
    """Pick the most distinctive ~10-word window from sentences spread over the document."""
    words = doc.tok.words
    wt = [0 if x in STOP else min(len(x), 10) for x in words]
    cands = []
    for (a, b) in sent_ranges:
        if b - a < win or any(excluded[a:b]):
            continue
        best_s, best_i = -1, a
        for i in range(a, b - win + 1):
            s = sum(wt[i:i + win])
            if s > best_s:
                best_s, best_i = s, i
        cands.append((a, best_s, best_i))
    if not cands:
        return []
    n = min(n, len(cands))
    size = len(cands) / n
    picks = []
    for bi in range(n):
        bucket = cands[int(bi * size): int((bi + 1) * size)] or [cands[min(int(bi * size), len(cands) - 1)]]
        _, _, i = max(bucket, key=lambda c: c[1])
        picks.append(" ".join(words[i:i + win]))
    return list(dict.fromkeys(picks))


# ============================================================================
# AI-content detection (local, free)
# ============================================================================
AI_PHRASES = [
    "delve", "tapestry", "testament to", "it is important to note", "it is worth noting",
    "in today's", "in the realm of", "navigate the complexities", "multifaceted",
    "plays a crucial role", "play a crucial role", "a myriad of", "underscores", "underscore the",
    "pivotal", "landscape of", "in conclusion", "furthermore", "moreover", "additionally",
    "fosters", "foster a", "leverage", "leveraging", "seamless", "holistic", "comprehensive understanding",
    "ever-evolving", "rapidly evolving", "paramount", "intricate", "nuanced", "embark", "shed light on",
    "sheds light", "in summary", "to summarize", "key takeaways", "unlock the", "as an ai",
]


def _clamp(x, lo=0.0, hi=1.0):
    return max(lo, min(hi, x))


def stylometric_signals(text: str) -> dict:
    """Weak, dependency-free signals: sentence-length burstiness + LLM-favoured phrasing."""
    sents = [text[a:b] for a, b in split_sentences(text)]
    lens = [len(WORD_RE.findall(s)) for s in sents if s.strip()]
    n_words = sum(lens)
    if len(lens) < 8 or n_words == 0:
        return {"score": None}
    mean = sum(lens) / len(lens)
    burst = pstdev(lens) / mean if mean else 0.0
    low = text.lower()
    density = sum(low.count(p) for p in AI_PHRASES) / n_words * 1000
    s_burst = _clamp((0.70 - burst) / 0.40)
    s_phr = _clamp(density / 8)
    return {"score": round(0.55 * s_burst + 0.45 * s_phr, 3), "burstiness": round(burst, 3),
            "ai_phrase_per_1k_words": round(density, 2)}


class NeuralAIDetector:
    """
    Binoculars-style zero-shot detector (Hans et al., 2024): ratio of an observer
    model's log-perplexity to the observer/performer cross-perplexity. LOWER score
    => more machine-like. Runs locally; requires torch + transformers.
    The default threshold is UNCALIBRATED - call calibrate() on your own samples.
    """

    def __init__(self, observer="Qwen/Qwen2.5-0.5B", performer="Qwen/Qwen2.5-0.5B-Instruct",
                 threshold: float = 0.90, scale: float = 0.06, max_tokens: int = 256):
        import torch
        from transformers import AutoModelForCausalLM, AutoTokenizer
        self.torch = torch
        self.dev = "cuda" if torch.cuda.is_available() else "cpu"
        self.tok = AutoTokenizer.from_pretrained(observer)
        self.obs = AutoModelForCausalLM.from_pretrained(observer).to(self.dev).eval()
        self.perf = AutoModelForCausalLM.from_pretrained(performer).to(self.dev).eval()
        self.threshold, self.scale, self.max_tokens = threshold, scale, max_tokens

    def raw_score(self, text: str) -> Optional[float]:
        t = self.torch
        enc = self.tok(text, return_tensors="pt", truncation=True, max_length=self.max_tokens).to(self.dev)
        if enc.input_ids.shape[1] < 24:
            return None
        with t.no_grad():
            lo = self.obs(**enc).logits[:, :-1].float()
            lp = self.perf(**enc).logits[:, :-1].float()
        labels = enc.input_ids[:, 1:]
        ppl = t.nn.functional.cross_entropy(lo.transpose(1, 2), labels, reduction="none").mean()
        xppl = -(t.softmax(lp, -1) * t.log_softmax(lo, -1)).sum(-1).mean()
        return float(ppl / xppl)

    def prob_ai(self, raw: float) -> float:
        return 1 / (1 + math.exp((raw - self.threshold) / self.scale))

    def calibrate(self, human_texts: list[str], ai_texts: list[str], target_fpr: float = 0.01):
        """Pick the threshold giving <= target_fpr on human samples (lower raw => AI)."""
        h = sorted(s for s in (self.raw_score(x) for x in human_texts) if s is not None)
        a = [s for s in (self.raw_score(x) for x in ai_texts) if s is not None]
        if not h:
            raise ValueError("no usable human samples")
        cut = h[max(0, int(len(h) * target_fpr) - 1)] if len(h) * target_fpr >= 1 else h[0] - 1e-6
        self.threshold = cut
        tpr = sum(x < cut for x in a) / len(a) if a else float("nan")
        return {"threshold": cut, "tpr": tpr, "fpr": sum(x < cut for x in h) / len(h)}

    def windows(self, text: str, sent_spans, min_words=60, max_words=200):
        out, start, count, first = [], None, 0, None
        for (a, b) in sent_spans:
            wc = len(WORD_RE.findall(text[a:b]))
            if first is None:
                first, count = a, 0
            count += wc
            last = b
            if count >= min_words:
                out.append((first, last))
                first = None
        if first is not None and count >= 30:
            out.append((first, last))
        return out


@dataclass
class AIResult:
    score: Optional[float]          # 0..1 likelihood signal (NOT a proof)
    label: str                      # likely_human | uncertain | likely_ai | insufficient_text
    confidence: str                 # low | medium
    method: str
    flagged_fraction: float = 0.0
    segments: list = field(default_factory=list)
    signals: dict = field(default_factory=dict)
    caveat: str = ("AI-content scores are probabilistic signals, not proof. Detectors can misjudge "
                   "non-native English, formulaic or heavily edited text. Never use as sole evidence.")


def detect_ai(text: str, sent_spans, neural: Optional[NeuralAIDetector]) -> AIResult:
    words = len(WORD_RE.findall(text))
    if words < 120:
        return AIResult(None, "insufficient_text", "low", "none", signals={"words": words})
    sty = stylometric_signals(text)
    segs, neural_scores, weights = [], [], []
    if neural:
        for (a, b) in neural.windows(text, sent_spans):
            raw = neural.raw_score(text[a:b])
            if raw is None:
                continue
            p = neural.prob_ai(raw)
            wc = len(WORD_RE.findall(text[a:b]))
            segs.append({"start": a, "end": b, "ai_score": round(p, 3), "binoculars": round(raw, 3)})
            neural_scores.append(p)
            weights.append(wc)
    parts, method = [], []
    if neural_scores:
        n_avg = sum(s * w for s, w in zip(neural_scores, weights)) / sum(weights)
        parts.append((0.7, n_avg))
        method.append("binoculars")
    if sty.get("score") is not None:
        parts.append((0.3 if neural_scores else 1.0, sty["score"]))
        method.append("stylometry")
    if not parts:
        return AIResult(None, "insufficient_text", "low", "none", signals=sty)
    tw = sum(w for w, _ in parts)
    score = sum(w * s for w, s in parts) / tw
    label = "likely_human" if score < 0.35 else "likely_ai" if score >= 0.65 else "uncertain"
    flagged = (sum(w for s, w in zip(neural_scores, weights) if s >= 0.7) / sum(weights)) if weights else 0.0
    return AIResult(round(score, 3), label, "medium" if neural_scores else "low", "+".join(method),
                    round(flagged, 3), segs, sty)


# ============================================================================
# Report model
# ============================================================================
@dataclass
class Match:
    start: int
    end: int
    text: str
    kind: str                       # exact | near | paraphrase | corpus
    similarity: float
    source_url: str
    source_title: str
    source_type: str
    highlight_url: Optional[str] = None
    quoted: bool = False


@dataclass
class SourceSummary:
    url: str
    title: str
    type: str
    matched_words: int
    percent_of_document: float


@dataclass
class Report:
    similarity_percent: float
    eligible_words: int
    matched_words: int
    matches: list
    sources: list
    ai: Optional[AIResult]
    warnings: list
    stats: dict
    text: str


# ============================================================================
# The detector
# ============================================================================
class PlagiarismDetector:
    def __init__(self, cfg: Optional[Config] = None, backends: Optional[list] = None,
                 neural_ai: Optional[NeuralAIDetector] = None, store: Optional[FingerprintStore] = None):
        self.cfg = cfg or Config()
        self.http = Http(self.cfg)
        self.cache = DiskCache(self.cfg.cache_dir, self.cfg.cache_ttl_days, self.cfg.cache_max_mb)
        self.backends = build_backends(self.http, self.cfg, cache=self.cache) if backends is None else backends
        self.neural = neural_ai
        self.store = store
        self.local_sources: list[Source] = []
        self._embedder = None

    # ---- corpora -----------------------------------------------------------
    def add_local_corpus(self, folder: str):
        for p in sorted(Path(folder).rglob("*")):
            if p.suffix.lower() in {".txt", ".md", ".pdf", ".docx", ".doc", ".html", ".htm"}:
                try:
                    self.local_sources.append(Source(url=p.resolve().as_uri(), title=p.name,
                                                     kind="local", text=extract_text(str(p))))
                except Exception as e:
                    log.warning("skip %s: %s", p, e)

    # ---- public API --------------------------------------------------------
    def analyze_file(self, path: str, **kw) -> Report:
        return self.analyze(extract_text(path, ocr=self.cfg.ocr), **kw)

    def analyze(self, text: str, web: bool = True, check_ai: bool = True, progress=None) -> Report:
        """progress: optional callback(stage: str, fraction: float) for UIs / job queues."""
        t0, cfg = time.time(), self.cfg
        p = progress or (lambda *_: None)
        p("Reading and cleaning text", 0.03)
        warnings: list[str] = []
        clean, st = normalize_text(text)
        if st["hidden_chars"] or st["homoglyphs"]:
            warnings.append(f"Found {st['hidden_chars']} hidden and {st['homoglyphs']} look-alike characters "
                            "(possible detector evasion). Text was normalized before analysis.")

        body = clean
        if cfg.exclude_references:
            refs = [m for m in _REF_RE.finditer(clean) if m.start() > len(clean) * 0.4]
            if refs:
                body = clean[:refs[-1].start()].rstrip()
        doc = Prepared(body, cfg.k)
        n = len(doc.tok.words)
        if n < cfg.k * 2:
            return Report(0.0, n, 0, [], [], None, ["Text too short to analyze."], {}, clean)

        sent_spans = split_sentences(body)
        sent_ranges = _sentence_token_ranges(doc, sent_spans)
        excluded = self._excluded_mask(doc)

        p("Searching the web and academic sources", 0.12)
        # 1. gather candidate sources
        sources: list[Source] = list(self.local_sources)
        n_queries = 0
        if web and self.backends:
            queries = plan_queries(doc, sent_ranges, excluded, cfg.max_queries)
            n_queries = len(queries)
            sources += self._discover(queries)
        p("Downloading candidate sources", 0.40)
        self._fetch_all(sources)

        p("Comparing your text with sources", 0.65)
        # 2. exact / near matching
        raws: list[Raw] = []
        preps: dict[int, Prepared] = {}
        for i, s in enumerate(sources):
            if len(s.text) < 100:
                continue
            sp = Prepared(normalize_text(s.text)[0], cfg.k)
            preps[i] = sp
            raws += find_spans(doc, sp, i, cfg)

        # 3. persistent fingerprint store (previous submissions, offline corpora)
        if self.store:
            for doc_id, positions in self.store.query(winnow(doc.hashes, cfg.window)).items():
                title, url, kind = self.store.meta(doc_id)
                sources.append(Source(url=url or f"store:{doc_id[:12]}", title=title or "Stored document",
                                      kind=kind))
                raws += find_corpus_spans(doc, positions, len(sources) - 1, cfg)

        # 4. paraphrase pass
        if cfg.detect_paraphrase and preps:
            covered = [False] * n
            for r in raws:
                for i in range(r.a, r.b):
                    covered[i] = True
            raws += paraphrase_pass(doc, sent_ranges, covered, excluded, sources, preps, cfg,
                                    self._get_embedder() if cfg.use_semantic else None)

        p("Scoring", 0.85)
        # 5. aggregate
        matches, summaries, matched, eligible = self._aggregate(doc, excluded, raws, sources, preps)

        p("Checking for AI-generated writing", 0.92)
        # 6. AI detection
        ai = detect_ai(body, sent_spans, self.neural) if check_ai else None
        if web and not self.backends:
            warnings.append("No search backends configured: only local/stored corpora were checked.")
        return Report(
            similarity_percent=round(100 * matched / eligible, 1) if eligible else 0.0,
            eligible_words=eligible, matched_words=matched, matches=matches, sources=summaries,
            ai=ai, warnings=warnings,
            stats={"seconds": round(time.time() - t0, 2), "queries": n_queries,
                   "sources_checked": len([s for s in sources if len(s.text) >= 100])},
            text=clean)

    # ---- internals ---------------------------------------------------------
    def _excluded_mask(self, doc: Prepared) -> list[bool]:
        """Tokens that don't count: quoted passages and in-text citations."""
        mask = [False] * len(doc.tok.words)
        st = doc.tok.starts
        pats = []
        if self.cfg.exclude_quotes:
            pats.append(_QUOTE_RE)
        if self.cfg.exclude_citations:
            pats.append(_CITE_RE)
        for pat in pats:
            for m in pat.finditer(doc.text):
                for i in range(bisect_left(st, m.start()), bisect_left(st, m.end())):
                    mask[i] = True
        return mask

    def _discover(self, queries: list[str]) -> list[Source]:
        cfg, pool = self.cfg, {}
        jobs = [(b, q) for q in queries for b in self.backends]

        def run(job):
            b, q = job
            key = f"search:{b.name}:{q}"
            hit = self.cache.get_json(key, ttl=cfg.search_cache_hours * 3600)
            if hit is not None:
                return [Source(**d) for d in hit]
            try:
                res = b.search(q, cfg.results_per_query)
            except Exception as e:
                log.debug("%s failed: %s", b.name, e)
                return []
            if res:                                  # never cache empties (may be an outage)
                self.cache.put_json(key, [asdict(x) for x in res])
            return res

        with ThreadPoolExecutor(max_workers=min(8, max(1, len(self.backends) * 2))) as ex:
            for results in ex.map(run, jobs):
                for s in results:
                    host = (urlparse(s.url).hostname or "").lower()
                    if any(host == d or host.endswith("." + d) for d in cfg.exclude_domains):
                        continue
                    if re.search(r"\.(jpe?g|png|gif|mp4|mp3|zip|svg)$", s.url, re.I):
                        continue
                    key = _canon(s.url)
                    if key in pool:
                        pool[key].votes += 1
                        pool[key].text = pool[key].text or s.text
                        pool[key].fetch_url = pool[key].fetch_url or s.fetch_url
                    else:
                        s.votes = 1
                        pool[key] = s
        ranked = sorted(pool.values(), key=lambda s: (-s.votes, not s.text))
        return ranked[:cfg.max_sources]

    def _fetch_all(self, sources: list[Source]):
        todo = [s for s in sources if s.kind not in ("local", "corpus") and (s.fetch_url or not s.text)]
        with ThreadPoolExecutor(max_workers=self.cfg.fetch_workers) as ex:
            list(ex.map(self._fetch_one, todo))

    def _fetch_one(self, s: Source):
        url = s.fetch_url or s.url
        cached = self.cache.get(url)
        if cached is not None:                       # "" = recent failure (negative cache)
            self._merge_text(s, cached)
            return
        got = self.http.get(url, robots=(s.kind == "web"))
        text = ""
        if got:
            data, ctype, enc = got
            try:
                if "pdf" in ctype.lower() or url.lower().endswith(".pdf") or data[:4] == b"%PDF":
                    text = pdf_to_text(data, max_pages=80, ocr=False)
                else:
                    text = html_to_text(data.decode(enc or "utf-8", "replace"))
            except Exception as e:
                log.debug("extract failed %s: %s", url, e)
        if len(text) > 200:
            self.cache.put(url, text)
            self._merge_text(s, text)
        else:
            self.cache.put(url, "")

    @staticmethod
    def _merge_text(s: Source, text: str):
        """Keep any pre-supplied abstract; add fetched full text after it."""
        if text and text not in s.text:
            s.text = (s.text + "\n\n" + text) if s.text else text

    def _get_embedder(self):
        if self._embedder is None:
            try:
                from sentence_transformers import SentenceTransformer
                self._embedder = SentenceTransformer("all-MiniLM-L6-v2")
            except ImportError:
                log.warning("sentence-transformers not installed; semantic matching disabled")
                self.cfg.use_semantic = False
        return self._embedder

    def _highlight_url(self, s: Source, sp: Optional[Prepared], src_a: Optional[int]) -> Optional[str]:
        """Browser text-fragment link that scrolls to and highlights the matched passage."""
        if s.kind not in ("web", "encyclopedia") or sp is None or src_a is None or s.url.lower().endswith(".pdf"):
            return None
        k = self.cfg.k
        t = sp.tok
        b = min(src_a + k, len(t.words)) - 1
        frag = sp.text[t.starts[src_a]:t.ends[b]]
        frag = re.sub(r"\s+", " ", frag)
        return s.url.split("#")[0] + "#:~:text=" + quote(frag, safe="").replace("-", "%2D")

    def _aggregate(self, doc, excluded, raws, sources, preps):
        n = len(doc.tok.words)
        tk = doc.tok
        eligible = n - sum(excluded)
        quoted, live = [], []
        for r in raws:
            (quoted if sum(excluded[r.a:r.b]) >= 0.8 * (r.b - r.a) else live).append(r)
        # boilerplate filter: a SHORT match found independently in many sources is a stock phrase
        per_tok: dict[int, set] = defaultdict(set)
        for r in live:
            for i in range(r.a, r.b):
                per_tok[i].add(r.src)

        def is_common(r: Raw) -> bool:
            if r.b - r.a > self.cfg.common_phrase_max_tokens:
                return False
            c = sorted(len(per_tok[i]) for i in range(r.a, r.b))
            return c[len(c) // 2] >= self.cfg.common_phrase_sources

        live = [r for r in live if not is_common(r)]
        by_src: dict[int, list[Raw]] = defaultdict(list)
        for r in live:
            by_src[r.src].append(r)
        order = sorted(by_src, key=lambda s: -sum(r.b - r.a for r in by_src[s]))

        owner = [-1] * n
        kept: list[Raw] = []
        for s in order:
            for r in sorted(by_src[s], key=lambda r: -(r.b - r.a)):
                new = [i for i in range(r.a, r.b) if owner[i] == -1 and not excluded[i]]
                if len(new) < max(4, 0.4 * (r.b - r.a)):
                    continue
                for i in new:
                    owner[i] = s
                kept.append(r)

        def to_match(r: Raw, is_quoted: bool) -> Match:
            s = sources[r.src]
            ca, cb = tk.starts[r.a], tk.ends[r.b - 1]
            return Match(ca, cb, doc.text[ca:cb], r.kind, r.sim, s.url, s.title, s.kind,
                         self._highlight_url(s, preps.get(r.src), r.src_a), is_quoted)

        matches = [to_match(r, False) for r in kept] + [to_match(r, True) for r in quoted]
        matches.sort(key=lambda m: (m.start, m.quoted))

        counts = Counter(o for o in owner if o != -1)
        summaries = [SourceSummary(sources[s].url, sources[s].title, sources[s].kind, c,
                                   round(100 * c / eligible, 1) if eligible else 0.0)
                     for s, c in counts.most_common()]
        return matches, summaries, sum(counts.values()), eligible


# ============================================================================
# HTML report (self-contained; light/dark; colour-coded highlights + source links)
# ============================================================================
_HTML = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>OpenPlag report</title>
<style>
:root{--bg:#fff;--fg:#1c1c1e;--mut:#6b6b70;--card:#f4f4f6;--bd:#dcdce0;--ml:84%;--hi:#f6b8b8;--mid:#f7e3a1;--lo:#d3ecd3}
@media (prefers-color-scheme:dark){:root{--bg:#161618;--fg:#ececf0;--mut:#9a9aa2;--card:#222226;--bd:#38383e;--ml:30%;--hi:#7a2d2d;--mid:#6b5a1f;--lo:#2b4a2b}}
body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.6 system-ui,-apple-system,Segoe UI,sans-serif}
main{max-width:880px;margin:0 auto;padding:24px 16px 64px}
h1{font-size:1.3rem;margin:0 0 12px}
.row{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:16px}
.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:12px 16px;flex:1 1 260px}
.big{font-size:2.2rem;font-weight:700;line-height:1.1}
.mut{color:var(--mut);font-size:.85rem}
ol{padding-left:0;list-style:none;margin:0}
ol li{display:flex;gap:10px;margin:8px 0;align-items:flex-start;overflow-wrap:anywhere}
.sw{flex:none;width:14px;height:14px;border-radius:3px;margin-top:5px;background:hsl(var(--h) 70% var(--ml))}
.tabs button{border:1px solid var(--bd);background:var(--card);color:var(--fg);padding:6px 14px;border-radius:8px;cursor:pointer;margin-right:6px}
.tabs button.on{border-color:var(--fg);font-weight:600}
.doc{white-space:pre-wrap;overflow-wrap:anywhere;border:1px solid var(--bd);border-radius:10px;padding:16px;margin-top:12px}
a.m,span.m{background:hsl(var(--h) 70% var(--ml));color:inherit;text-decoration:none;border-radius:3px}
a.m:hover{outline:2px solid var(--fg)}
.m.q{background:none;border-bottom:2px dashed var(--mut)}
.m.paraphrase{border-bottom:2px dotted var(--fg)}
sup{font-size:.7em;margin-left:1px;font-weight:700}
.ai.hi{background:var(--hi)}.ai.mid{background:var(--mid)}.ai.lo{background:var(--lo)}
.warn{border-left:4px solid #d9a400;padding:6px 12px;margin:8px 0;background:var(--card)}
a{color:inherit}
</style></head><body><main>
<h1>OpenPlag similarity report</h1>
<div class="row">
 <div class="card"><div class="big">@@PCT@@%</div><div class="mut">similarity - @@MATCHED@@ of @@ELIGIBLE@@ counted words matched</div></div>
 @@AICARD@@
</div>
@@WARNINGS@@
<div class="card" style="margin-bottom:16px"><b>Sources</b><ol>@@SOURCES@@</ol></div>
<div class="tabs"><button id="bp" class="on">Plagiarism view</button><button id="ba">AI-content view</button></div>
<div class="doc" id="vp">@@PLAG@@</div>
<div class="doc" id="va" hidden>@@AIVIEW@@</div>
<p class="mut">Solid colour = near-verbatim copy, dotted underline = paraphrase, dashed = quote/citation (not counted).
Click a highlight to open the source (supported browsers jump to the matching passage).</p>
</main>
<script>
const P=document.getElementById('vp'),A=document.getElementById('va'),bp=document.getElementById('bp'),ba=document.getElementById('ba');
bp.onclick=()=>{P.hidden=false;A.hidden=true;bp.className='on';ba.className=''};
ba.onclick=()=>{P.hidden=true;A.hidden=false;ba.className='on';bp.className=''};
</script></body></html>"""


def report_from_dict(d: dict) -> Report:
    """Rebuild a Report from asdict(report) / stored JSON (used by the API to re-render HTML)."""
    ai = AIResult(**d["ai"]) if d.get("ai") else None
    return Report(similarity_percent=d["similarity_percent"], eligible_words=d["eligible_words"],
                  matched_words=d["matched_words"], matches=[Match(**m) for m in d["matches"]],
                  sources=[SourceSummary(**x) for x in d["sources"]], ai=ai,
                  warnings=d.get("warnings", []), stats=d.get("stats", {}), text=d["text"])


def _safe_href(u: str) -> Optional[str]:
    return u if urlparse(u).scheme in ("http", "https") else None


def render_html(rep: Report) -> str:
    esc = htmllib.escape
    text = rep.text
    idx = {s.url: i + 1 for i, s in enumerate(rep.sources)}
    hue = lambda i: (i * 47) % 360

    out, cur = [], 0
    for m in sorted(rep.matches, key=lambda m: (m.start, m.quoted)):
        a = max(m.start, cur)
        if m.end <= a:
            continue
        out.append(esc(text[cur:a]))
        n = idx.get(m.source_url)
        label = "excluded" if m.quoted else m.kind
        tip = f"{label} match ({m.similarity:.0%}) - {m.source_title or m.source_url}"
        href = _safe_href(m.highlight_url or m.source_url)
        tag = "a" if href else "span"
        attrs = f' href="{esc(href, quote=True)}" target="_blank" rel="noopener noreferrer"' if href else ""
        sup = f"<sup>{n}</sup>" if n and not m.quoted else ""
        cls = "q" if m.quoted else m.kind
        out.append(f'<{tag} class="m {cls}" style="--h:{hue(n or 0)}"{attrs} title="{esc(tip, quote=True)}">'
                   f'{esc(text[a:m.end])}{sup}</{tag}>')
        cur = m.end
    out.append(esc(text[cur:]))

    segs = sorted(rep.ai.segments, key=lambda x: x["start"]) if rep.ai and rep.ai.segments else []
    if segs:
        o, cur = [], 0
        for sg in segs:
            a = max(sg["start"], cur)
            if sg["end"] <= a:
                continue
            o.append(esc(text[cur:a]))
            c = "hi" if sg["ai_score"] >= 0.65 else "mid" if sg["ai_score"] >= 0.35 else "lo"
            o.append(f'<span class="ai {c}" title="AI signal {sg["ai_score"]:.0%}">{esc(text[a:sg["end"]])}</span>')
            cur = sg["end"]
        o.append(esc(text[cur:]))
        ai_view = "".join(o)
    else:
        ai_view = ('<span class="mut">Per-passage AI highlighting needs the neural detector '
                   '(--neural-ai). Showing plain text.</span>\n\n' + esc(text))

    items = []
    for i, sm in enumerate(rep.sources, 1):
        href = _safe_href(sm.url)
        name = esc(sm.title or sm.url)
        link = f'<a href="{esc(href, quote=True)}" target="_blank" rel="noopener noreferrer">{name}</a>' if href else name
        items.append(f'<li style="--h:{hue(i)}"><span class="sw"></span><div>{i}. {link}'
                     f'<div class="mut">{esc(sm.type)} - {sm.percent_of_document}% of document - '
                     f'{sm.matched_words} words<br>{esc(sm.url)}</div></div></li>')
    if not items:
        items = ['<li class="mut">No matching sources found.</li>']

    ai = rep.ai
    ai_card = ""
    if ai:
        score = "n/a" if ai.score is None else f"{ai.score:.0%}"
        ai_card = (f'<div class="card"><div class="big">{score}</div><div class="mut">AI-content signal: '
                   f'{esc(ai.label.replace("_", " "))} (confidence {esc(ai.confidence)}; {esc(ai.method)})<br>'
                   f'{esc(ai.caveat)}</div></div>')
    warns = "".join(f'<div class="warn">{esc(w)}</div>' for w in rep.warnings)

    html = _HTML
    for k, v in {"@@PCT@@": str(rep.similarity_percent), "@@MATCHED@@": str(rep.matched_words),
                 "@@ELIGIBLE@@": str(rep.eligible_words), "@@AICARD@@": ai_card, "@@WARNINGS@@": warns,
                 "@@SOURCES@@": "".join(items), "@@PLAG@@": "".join(out), "@@AIVIEW@@": ai_view}.items():
        html = html.replace(k, v)
    return html


# ============================================================================
# CLI
# ============================================================================
def _print_report(rep: Report):
    print(f"\nSimilarity: {rep.similarity_percent}%  ({rep.matched_words}/{rep.eligible_words} words)")
    for w in rep.warnings:
        print("  !", w)
    print("\nTop sources:")
    for sm in rep.sources[:10]:
        print(f"  {sm.percent_of_document:5.1f}%  [{sm.type}] {sm.title[:60]}\n          {sm.url}")
    print("\nMatches:")
    for m in rep.matches[:15]:
        tag = "excluded" if m.quoted else m.kind
        print(f"  - ({tag}, {m.similarity:.2f}) \"{m.text[:90]}...\"\n      {m.highlight_url or m.source_url}")
    if rep.ai:
        a = rep.ai
        print(f"\nAI-content signal: {a.label} (score={a.score}, confidence={a.confidence}, via {a.method})")
        if a.flagged_fraction:
            print(f"  ~{a.flagged_fraction:.0%} of text in high-scoring segments")
        print("  Note:", a.caveat)
    print(f"\n{rep.stats}")


def main():
    ap = argparse.ArgumentParser(description="OpenPlag: free plagiarism + AI-content detection")
    ap.add_argument("input", nargs="?", help="file path (.pdf .docx .doc .txt) or '-' for stdin")
    ap.add_argument("--text", help="analyze pasted text instead of a file")
    ap.add_argument("--json", help="write full JSON report here")
    ap.add_argument("--html", help="write a self-contained HTML report with highlights here")
    ap.add_argument("--local-corpus", help="folder of reference documents to compare against")
    ap.add_argument("--store", help="SQLite fingerprint store path (previous submissions)")
    ap.add_argument("--add-to-store", action="store_true", help="after analysis, add this document to --store")
    ap.add_argument("--searxng", default=os.environ.get("OPENPLAG_SEARXNG", ""),
                    help="base URL of YOUR SearXNG instance, e.g. http://localhost:8080 (env OPENPLAG_SEARXNG)")
    ap.add_argument("--searx-instances", default="", help="comma list of extra SearXNG instances (failover order)")
    ap.add_argument("--public-searx", action="store_true",
                    help="also try public SearXNG instances from searx.space (best-effort; most block API use)")
    ap.add_argument("--check-searx", action="store_true", help="probe SearXNG instances for JSON support, then exit")
    ap.add_argument("--backends", default=",".join(DEFAULT_BACKENDS), help="comma list: " + ",".join(BACKENDS))
    ap.add_argument("--no-web", action="store_true")
    ap.add_argument("--no-ai", action="store_true")
    ap.add_argument("--no-ocr", action="store_true")
    ap.add_argument("--neural-ai", action="store_true", help="enable Binoculars (torch+transformers)")
    ap.add_argument("--semantic", action="store_true", help="semantic paraphrase matching")
    ap.add_argument("--email", default="", help="contact email for polite API pools")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args()
    logging.basicConfig(level=logging.DEBUG if a.verbose else logging.WARNING, format="%(message)s")

    cfg = Config(use_semantic=a.semantic, contact_email=a.email, ocr=not a.no_ocr)
    if a.searx_instances:
        cfg.searx_instances = tuple(u.strip() for u in a.searx_instances.split(",") if u.strip())

    if a.check_searx:
        http = Http(cfg)
        pool = SearXNGPool(http, cfg, ([a.searxng] if a.searxng else []) + list(cfg.searx_instances),
                           discover=cfg.searx_discover, cache=DiskCache(cfg.cache_dir, cfg.cache_ttl_days))
        pool._try_discover()
        print(f"{'instance':45} {'JSON':5} {'reason':15} results   ms")
        for r in pool.probe():
            print(f"{r['instance']:45} {'yes' if r['ok'] else 'no':5} {r['reason']:15} {r['results']:7} {r['ms']:5}")
        return

    if a.text:
        text = a.text
    elif a.input == "-":
        text = sys.stdin.read()
    elif a.input:
        text = extract_text(a.input, ocr=cfg.ocr)
    else:
        ap.error("provide a file, '-', --text, or --check-searx")

    store = FingerprintStore(a.store, cfg) if a.store else None
    det = PlagiarismDetector(cfg, backends=[], store=store,
                             neural_ai=NeuralAIDetector() if a.neural_ai and not a.no_ai else None)
    det.backends = build_backends(det.http, cfg, [n for n in a.backends.split(",") if n], a.searxng,
                                  public_searx=a.public_searx, cache=det.cache)
    if a.local_corpus:
        det.add_local_corpus(a.local_corpus)
    rep = det.analyze(text, web=not a.no_web, check_ai=not a.no_ai)
    _print_report(rep)
    if a.json:
        Path(a.json).write_text(json.dumps(asdict(rep), indent=2, ensure_ascii=False))
        print("JSON report ->", a.json)
    if a.html:
        Path(a.html).write_text(render_html(rep), encoding="utf-8")
        print("HTML report ->", a.html)
    if store and a.add_to_store:
        store.add(text, title=a.input or "pasted text")


if __name__ == "__main__":
    main()
